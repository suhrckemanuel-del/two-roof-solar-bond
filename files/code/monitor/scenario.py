"""Market scenario on a COPY of the workbook: write the new inputs with Excel COM, recalculate, export facts.

The base workbook and every model/*facts.json stay untouched. Facts are exported by running the project's own
scripts (export_facts.py, proposal_facts.py, feasibility_facts.py) unchanged, in a scratch folder that holds a copy of
the scenario workbook under the name the scripts expect. Their outputs are moved into <run>/facts/.
"""
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

MODEL = Path(__file__).resolve().parents[1]
BASE_XLSX = MODEL / "solar_bond_model.xlsx"
FACT_SCRIPTS = ("export_facts.py", "proposal_facts.py", "feasibility_facts.py")
UNCHANGED = ("mc_facts.json",)  # Monte Carlo is not re-run; its base facts are carried over as they are


def input_rows(path):
    ws = load_workbook(path, data_only=True)["Inputs"]
    return {ws.cell(r, 1).value: (r, ws.cell(r, 3).value) for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value}


def recalc_copy(dst, changes, log=print):
    """Copy the base workbook to dst, write {input_key: value} into Inputs!C<row> of the copy, CalculateFull, save."""
    shutil.copy2(BASE_XLSX, dst)
    rows = input_rows(dst)
    missing = [k for k in changes if k not in rows]
    if missing:
        raise KeyError(f"inputs not on the Inputs sheet: {missing}")
    import pythoncom
    import win32com.client as w
    pythoncom.CoInitialize()
    xl = w.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(str(Path(dst).resolve()))
        ws = wb.Worksheets("Inputs")
        for k, v in changes.items():
            r = rows[k][0]
            if ws.Cells(r, 3).HasFormula:
                raise ValueError(f"Inputs!C{r} ({k}) holds a formula; refusing to overwrite it")
            ws.Cells(r, 3).Value = v
            log(f"  wrote Inputs!C{r} {k} = {v} (was {rows[k][1]})")
        xl.CalculateFull()
        wb.Save()
        wb.Close(True)
    finally:
        xl.Quit()
    # read back what Excel saved: the copy must hold exactly the values we wrote
    after = input_rows(dst)
    for k, v in changes.items():
        if abs(float(after[k][1]) - float(v)) > 1e-9:
            raise RuntimeError(f"read-back mismatch for {k}: {after[k][1]} != {v}")
    return {k: {"cell": f"Inputs!C{rows[k][0]}", "old": rows[k][1], "new": v} for k, v in changes.items()}


def export_facts(xlsx, out_dir, log=print):
    """Run the project's fact scripts against xlsx; write facts.json + *_facts.json into out_dir."""
    out_dir = Path(out_dir)
    work = out_dir.parent / "_work"
    if work.exists():
        shutil.rmtree(work)
    work.mkdir(parents=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copy2(xlsx, work / "solar_bond_model.xlsx")
        for s in ("mc_engine.py",) + FACT_SCRIPTS:
            shutil.copy2(MODEL / s, work / s)
        for s in FACT_SCRIPTS:
            r = subprocess.run([sys.executable, s], cwd=work, capture_output=True, text=True)
            if r.returncode:
                raise RuntimeError(f"{s} failed: {r.stderr[-800:]}")
            log(f"  {s}: {r.stdout.strip().splitlines()[0].split(' to ')[0] if s == 'export_facts.py' else 'ok'}")
        for f in ("facts.json", "proposal_facts.json", "feasibility_facts.json"):
            shutil.move(str(work / f), out_dir / f)
        for f in UNCHANGED:
            shutil.copy2(MODEL / f, out_dir / f)
    finally:
        shutil.rmtree(work, ignore_errors=True)
    return out_dir
